"""Convert an Excel workbook (.xlsx/.xlsm) into the app's workbook JSON.

Usage: python tools/xlsx_to_workbook.py <input.xlsm> <output.json> [workbook name]

Keeps values, formulas, number formats, fonts, fills, borders, alignment,
merges, column widths, row heights, hidden sheets/rows/cols, frozen panes,
Excel tables and conditional formatting (cellIs / expression rules).
"""
import datetime
import json
import re
import sys
import zipfile

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

SRC = sys.argv[1]
OUT = sys.argv[2]
NAME = sys.argv[3] if len(sys.argv) > 3 else 'Reunion 123'

EXCEL_EPOCH = datetime.datetime(1899, 12, 30)

# ---------------------------------------------------------------- theme colors
with zipfile.ZipFile(SRC) as z:
    theme_xml = z.read('xl/theme/theme1.xml').decode('utf8')
scheme = re.search(r'<a:clrScheme.*?</a:clrScheme>', theme_xml, re.S).group(0)
raw = []
for tag in ['dk1', 'lt1', 'dk2', 'lt2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink']:
    block = re.search(rf'<a:{tag}>(.*?)</a:{tag}>', scheme, re.S).group(1)
    m = re.search(r'lastClr="([0-9A-Fa-f]{6})"', block) or re.search(r'val="([0-9A-Fa-f]{6})"', block)
    raw.append(m.group(1))
THEME = [raw[1], raw[0], raw[3], raw[2]] + raw[4:]


def apply_tint(hex6, tint):
    r, g, b = (int(hex6[i:i + 2], 16) for i in (0, 2, 4))
    if tint:
        f = (lambda c: c + (255 - c) * tint) if tint > 0 else (lambda c: c * (1 + tint))
        r, g, b = (max(0, min(255, round(f(c)))) for c in (r, g, b))
    return f'#{r:02x}{g:02x}{b:02x}'


def resolve_color(col):
    if col is None:
        return None
    try:
        if col.type == 'rgb' and isinstance(col.rgb, str) and len(col.rgb) == 8:
            if col.rgb == '00000000':
                return None
            return apply_tint(col.rgb[2:], col.tint)
        if col.type == 'theme' and col.theme is not None and col.theme < len(THEME):
            return apply_tint(THEME[col.theme], col.tint)
        if col.type == 'indexed' and col.indexed is not None and col.indexed < len(COLOR_INDEX):
            if col.indexed in (64, 65):  # system foreground/background
                return None
            return apply_tint(COLOR_INDEX[col.indexed][2:], col.tint)
    except Exception:
        pass
    return None


def to_serial(v):
    if isinstance(v, datetime.datetime):
        d = v - EXCEL_EPOCH
    elif isinstance(v, datetime.date):
        d = datetime.datetime(v.year, v.month, v.day) - EXCEL_EPOCH
    elif isinstance(v, datetime.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    else:
        return None
    serial = d.days + d.seconds / 86400
    return int(serial) if float(serial).is_integer() else round(serial, 10)


def font_style(f, s):
    if f is None:
        return
    if f.name:
        s['fontFamily'] = f.name
    if f.sz and abs(f.sz - 11) > 0.01:
        s['fontSize'] = round(f.sz, 1)
    if f.b: s['bold'] = True
    if f.i: s['italic'] = True
    if f.u: s['underline'] = True
    if f.strike: s['strike'] = True
    fc = resolve_color(f.color)
    if fc and fc != '#000000':
        s['color'] = fc


def cell_style(c):
    s = {}
    font_style(c.font, s)
    if c.fill and c.fill.fill_type == 'solid':
        bg = resolve_color(c.fill.fgColor)
        if bg and bg != '#ffffff':
            s['bg'] = bg
    a = c.alignment
    if a.horizontal in ('left', 'center', 'right'):
        s['align'] = a.horizontal
    elif a.horizontal in ('centerContinuous', 'distributed', 'justify'):
        s['align'] = 'center'
    if a.vertical in ('top', 'center', 'bottom'):
        s['valign'] = 'middle' if a.vertical == 'center' else a.vertical
    if a.wrap_text:
        s['wrap'] = True
    if a.indent:
        s['indent'] = int(a.indent)
    bd = {}
    for k, side in (('top', c.border.top), ('right', c.border.right), ('bottom', c.border.bottom), ('left', c.border.left)):
        if side is not None and side.style:
            bd[k] = True
    if bd:
        s['border'] = bd
    nf = c.number_format
    if nf and nf != 'General':
        s['nf'] = nf
    return s


def dxf_style(dxf):
    s = {}
    if dxf is None:
        return s
    if dxf.font is not None:
        fc = resolve_color(dxf.font.color)
        if fc: s['color'] = fc
        if dxf.font.b: s['bold'] = True
        if dxf.font.i: s['italic'] = True
    if dxf.fill is not None:
        bg = resolve_color(dxf.fill.bgColor) or resolve_color(dxf.fill.fgColor)
        if bg: s['bg'] = bg
    return s


import os
import posixpath

ASSET_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'assets')
EMU_PER_PX = 9525


def sheet_images():
    """Pictures per sheet title: position (col/row + px offsets), size and assigned macro."""
    out = {}
    with zipfile.ZipFile(SRC) as z:
        names = set(z.namelist())
        wbx = z.read('xl/workbook.xml').decode('utf8')
        rels = z.read('xl/_rels/workbook.xml.rels').decode('utf8')
        rid_target = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
        rid_target.update({a: b for b, a in re.findall(r'Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
        for title, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
            target = rid_target.get(rid)
            if not target:
                continue
            sheet_path = posixpath.normpath(posixpath.join('xl', target.lstrip('/').replace('xl/', '', 1)))
            rel_path = posixpath.join(posixpath.dirname(sheet_path), '_rels', posixpath.basename(sheet_path) + '.rels')
            if rel_path not in names:
                continue
            srels = z.read(rel_path).decode('utf8')
            for dtarget in re.findall(r'Target="([^"]*drawings/[^"]+)"', srels):
                dpath = posixpath.normpath(posixpath.join(posixpath.dirname(sheet_path), dtarget))
                dxml = z.read(dpath).decode('utf8')
                drel_path = posixpath.join(posixpath.dirname(dpath), '_rels', posixpath.basename(dpath) + '.rels')
                drels = z.read(drel_path).decode('utf8') if drel_path in names else ''
                media = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', drels))
                for anchor in re.findall(r'<xdr:(?:twoCellAnchor|oneCellAnchor)[\s\S]*?</xdr:(?:twoCellAnchor|oneCellAnchor)>', dxml):
                    pic = re.search(r'<xdr:pic(?: macro="([^"]*)")?', anchor)
                    blip = re.search(r'r:embed="([^"]+)"', anchor)
                    if not pic or not blip:
                        continue
                    fr = re.search(r'<xdr:from><xdr:col>(\d+)</xdr:col><xdr:colOff>(-?\d+)</xdr:colOff><xdr:row>(\d+)</xdr:row><xdr:rowOff>(-?\d+)</xdr:rowOff>', anchor)
                    to = re.search(r'<xdr:to><xdr:col>(\d+)</xdr:col><xdr:colOff>(-?\d+)</xdr:colOff><xdr:row>(\d+)</xdr:row><xdr:rowOff>(-?\d+)</xdr:rowOff>', anchor)
                    ext = re.search(r'<a:ext cx="(\d+)" cy="(\d+)"', anchor)
                    mpath = posixpath.normpath(posixpath.join(posixpath.dirname(dpath), media.get(blip.group(1), '')))
                    if mpath not in names:
                        continue
                    os.makedirs(ASSET_DIR, exist_ok=True)
                    fname = f"{re.sub(r'[^A-Za-z0-9_-]', '_', title)}_{posixpath.basename(mpath)}"
                    with open(os.path.join(ASSET_DIR, fname), 'wb') as fh:
                        fh.write(z.read(mpath))
                    img = {
                        'src': f'assets/{fname}',
                        'col': int(fr.group(1)) + 1, 'dx': round(int(fr.group(2)) / EMU_PER_PX),
                        'row': int(fr.group(3)) + 1, 'dy': round(int(fr.group(4)) / EMU_PER_PX),
                    }
                    if ext:
                        img['w'] = round(int(ext.group(1)) / EMU_PER_PX)
                        img['h'] = round(int(ext.group(2)) / EMU_PER_PX)
                    if to:
                        img['toCol'] = int(to.group(1)) + 1
                        img['toDx'] = round(int(to.group(2)) / EMU_PER_PX)
                    if pic.group(1):
                        img['macro'] = pic.group(1).split('!')[-1]
                    out.setdefault(title, []).append(img)
    return out


IMAGES = sheet_images()


def read_pivots():
    """Pivot tables: source range, row/column/value/filter fields and hidden items."""
    import xml.etree.ElementTree as ET
    NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    R_NS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    AGG = {None: 'sum', 'sum': 'sum', 'count': 'count', 'average': 'average', 'max': 'max', 'min': 'min',
           'product': 'product', 'countNums': 'countNums', 'stdDev': 'stdDev', 'var': 'var'}
    out = []
    with zipfile.ZipFile(SRC) as z:
        names = set(z.namelist())
        wbx = z.read('xl/workbook.xml').decode('utf8')
        rels = z.read('xl/_rels/workbook.xml.rels').decode('utf8')
        rid_target = {}
        for m in re.finditer(r'<Relationship [^>]*>', rels):
            tag = m.group(0)
            i = re.search(r'Id="([^"]+)"', tag); t = re.search(r'Target="([^"]+)"', tag)
            if i and t: rid_target[i.group(1)] = t.group(1)
        for title, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
            title = title.replace('&amp;', '&').replace('&apos;', "'").replace('&quot;', '"')
            target = rid_target.get(rid)
            if not target:
                continue
            sheet_path = posixpath.normpath(posixpath.join('xl', target.lstrip('/').replace('xl/', '', 1)))
            rel_path = posixpath.join(posixpath.dirname(sheet_path), '_rels', posixpath.basename(sheet_path) + '.rels')
            if rel_path not in names:
                continue
            srels = z.read(rel_path).decode('utf8')
            for ptarget in re.findall(r'Target="([^"]*pivotTables/[^"]+)"', srels):
                ppath = posixpath.normpath(posixpath.join(posixpath.dirname(sheet_path), ptarget))
                prel = posixpath.join(posixpath.dirname(ppath), '_rels', posixpath.basename(ppath) + '.rels')
                ctarget = re.search(r'Target="([^"]+)"', z.read(prel).decode('utf8')).group(1)
                cpath = posixpath.normpath(posixpath.join(posixpath.dirname(ppath), ctarget))
                pt = ET.fromstring(z.read(ppath))
                cd = ET.fromstring(z.read(cpath))
                ws_src = cd.find('m:cacheSource/m:worksheetSource', NS)
                if ws_src is None:
                    continue
                fields = []
                for cf in cd.findall('m:cacheFields/m:cacheField', NS):
                    items = []
                    si = cf.find('m:sharedItems', NS)
                    if si is not None:
                        for it in si:
                            tag = it.tag.split('}')[1]
                            v = it.get('v')
                            if tag == 'n': v = float(v); v = int(v) if v.is_integer() else v
                            elif tag == 'm': v = None
                            elif tag == 'b': v = v in ('1', 'true')
                            elif tag == 'd':
                                dt = datetime.datetime.fromisoformat(v)
                                v = to_serial(dt)
                            items.append(v)
                    fields.append({'name': cf.get('name'), 'items': items})
                fname = lambda i: fields[i]['name'] if 0 <= i < len(fields) else None
                hidden, collapsed, order, sort = {}, {}, {}, {}
                layout = 'compact' if pt.get('compact') != '0' else ('outline' if pt.get('outline') == '1' else 'tabular')
                for idx, pf in enumerate(pt.findall('m:pivotFields/m:pivotField', NS)):
                    hid, col, ordr = [], [], []
                    its = fields[idx]['items'] if idx < len(fields) else []
                    for it in pf.findall('m:items/m:item', NS):
                        if it.get('x') is None or it.get('t'):
                            continue
                        x = int(it.get('x'))
                        if x >= len(its):
                            continue
                        ordr.append(its[x])
                        if it.get('h') == '1': hid.append(its[x])
                        if it.get('sd') == '0': col.append(its[x])
                    if hid: hidden[fname(idx)] = hid
                    if col: collapsed[fname(idx)] = col
                    if pf.get('axis') in ('axisRow', 'axisCol'):
                        if pf.get('sortType') in ('ascending', 'descending'): sort[fname(idx)] = pf.get('sortType')
                        elif ordr: order[fname(idx)] = ordr
                rows = [fname(int(f.get('x'))) if int(f.get('x')) >= 0 else '__values' for f in pt.findall('m:rowFields/m:field', NS)]
                cols = [fname(int(f.get('x'))) if int(f.get('x')) >= 0 else '__values' for f in pt.findall('m:colFields/m:field', NS)]
                values = [{'field': fname(int(df.get('fld'))), 'agg': AGG.get(df.get('subtotal'), 'sum'), 'name': df.get('name')}
                          for df in pt.findall('m:dataFields/m:dataField', NS)]
                filters = []
                for pf in pt.findall('m:pageFields/m:pageField', NS):
                    fld = int(pf.get('fld'))
                    item = pf.get('item')
                    sel = None
                    if item is not None and fld < len(fields) and int(item) < len(fields[fld]['items']):
                        sel = [fields[fld]['items'][int(item)]]
                    filters.append({'field': fname(fld), 'selected': sel})
                loc = pt.find('m:location', NS)
                src_sheet = ws_src.get('sheet') or title
                out.append({
                    'name': pt.get('name'),
                    'sheet': title,
                    'location': loc.get('ref'),
                    'firstHeaderRow': int(loc.get('firstHeaderRow', 1)),
                    'firstDataRow': int(loc.get('firstDataRow', 1)),
                    'firstDataCol': int(loc.get('firstDataCol', 1)),
                    'source': {'sheet': src_sheet, 'ref': ws_src.get('ref') or ws_src.get('name')},
                    'rows': rows, 'cols': cols, 'values': values, 'filters': filters, 'hidden': hidden,
                    'grandRows': pt.get('rowGrandTotals', '1') != '0',
                    'grandCols': pt.get('colGrandTotals', '1') != '0',
                    'layout': layout, 'collapsed': collapsed, 'order': order, 'sort': sort,
                    'imported': True
                })
    return out


PIVOTS = read_pivots()
print(f'Tablas dinámicas: {len(PIVOTS)}', flush=True)

print('Loading…', flush=True)
wb = openpyxl.load_workbook(SRC, keep_vba=True)
book = {'workbookName': NAME, 'sheetOrder': [], 'sheets': {}, 'tables': {}}
stats = {}

def data_bounds(ws):
    """Last row/col that hold a value (formatting-only cells beyond this are dropped)."""
    max_r = max_c = 0
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and str(c.value) != '':
                if c.row > max_r: max_r = c.row
                if c.column > max_c: max_c = c.column
    return max_r, max_c


for ws in wb.worksheets:
    cells = {}
    n_formulas = 0
    last_r, last_c = data_bounds(ws)
    # formatting-only cells outside the data area are kept when they draw something
    # (border / fill); whole-column formats (1M rows) are cut a bit past the data
    scan_r = ws.max_row if ws.max_row <= 20000 else last_r + 500
    scan_c = ws.max_column if ws.max_column <= 300 else last_c + 20
    for row in ws.iter_rows(max_row=max(scan_r, 1), max_col=max(scan_c, 1)):
        for c in row:
            v = c.value
            outside = c.row > last_r or c.column > last_c
            if outside and (v is None or str(v) == ''):
                s_out = cell_style(c)
                if 'border' in s_out or 'bg' in s_out:
                    cells[c.coordinate] = {'s': s_out}
                continue
            s = cell_style(c)
            entry = {}
            if isinstance(v, ArrayFormula):
                v = v.text
            if isinstance(v, str) and v.startswith('='):
                entry = {'f': v, 'v': ''}
                n_formulas += 1
            elif isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
                entry = {'v': to_serial(v)}
                if 'nf' not in s:
                    s['nf'] = 'm/d/yyyy'
            elif isinstance(v, bool):
                entry = {'v': v}
            elif isinstance(v, (int, float)):
                entry = {'v': int(v) if float(v).is_integer() and abs(v) < 1e15 else v}
            elif v is not None and str(v) != '':
                entry = {'v': str(v)}
            meaningful = {k: x for k, x in s.items() if k not in ('fontFamily',)}
            if entry or meaningful:
                if s:
                    entry['s'] = s
                cells[c.coordinate] = entry

    col_widths, hidden_cols = {}, []
    for key, dim in ws.column_dimensions.items():
        for i in range(dim.min or 0, (dim.max or 0) + 1):
            if i < 1 or i > 16384:
                continue
            letter = get_column_letter(i)
            if dim.width:
                col_widths[letter] = round(dim.width, 2)
            if dim.hidden:
                hidden_cols.append(letter)

    row_heights, hidden_rows = {}, []
    for r, dim in ws.row_dimensions.items():
        if dim.height:
            row_heights[str(r)] = round(dim.height, 2)
        if dim.hidden:
            hidden_rows.append(r)

    freeze = None
    pane = ws.sheet_view.pane
    if pane is not None and pane.state in ('frozen', 'frozenSplit'):
        freeze = {'rows': int(pane.ySplit or 0), 'cols': int(pane.xSplit or 0)}

    cond = []
    for cf in ws.conditional_formatting:
        rules = []
        for rule in cf.rules:
            if rule.type not in ('cellIs', 'expression'):
                print(f'  [aviso] {ws.title}: regla condicional no soportada: {rule.type}')
                continue
            rules.append({
                'type': rule.type,
                'operator': rule.operator,
                'formulas': list(rule.formula or []),
                'priority': rule.priority,
                'stopIfTrue': bool(rule.stopIfTrue),
                'style': dxf_style(rule.dxf)
            })
        if rules:
            cond.append({'ranges': str(cf.sqref).split(), 'rules': rules})

    def read_filter(af, ref, table_name=None):
        cols = {}
        for fc in (af.filterColumn or []):
            spec = {}
            if fc.filters is not None:
                spec['values'] = [str(f) for f in (fc.filters.filter or [])]
                if fc.filters.blank:
                    spec['blanks'] = True
            if fc.customFilters is not None:
                spec['custom'] = [{'op': cf.operator or 'equal', 'val': cf.val} for cf in fc.customFilters.customFilter]
                spec['and'] = bool(fc.customFilters._and)
            cols[str(fc.colId)] = spec
        out = {'ref': ref, 'columns': cols}
        if table_name:
            out['table'] = table_name
        return out

    filters = []
    for table in ws.tables.values():
        book['tables'][table.name] = {
            'sheet': ws.title,
            'ref': table.ref,
            'columns': [col.name for col in table.tableColumns]
        }
        if table.autoFilter is not None:
            filters.append(read_filter(table.autoFilter, table.autoFilter.ref or table.ref, table.name))
    if ws.auto_filter is not None and ws.auto_filter.ref and ':' in ws.auto_filter.ref:
        filters.append(read_filter(ws.auto_filter, ws.auto_filter.ref))
    # rows hidden inside a filter range are kept apart: the app recomputes them when a filter is applied
    filtered = []
    if filters and hidden_rows:
        def in_filter(r):
            for f in filters:
                a, b = f['ref'].split(':')
                r1 = int(re.sub(r'[A-Z$]', '', a)); r2 = int(re.sub(r'[A-Z$]', '', b))
                if r1 < r <= r2:
                    return True
            return False
        filtered = [r for r in hidden_rows if in_filter(r)]
        hidden_rows = [r for r in hidden_rows if not in_filter(r)]
        print(f'  {ws.title}: filtros={len(filters)}, filas ocultas por filtro={len(filtered)}')

    # sheet-wide default style = the most common one; cells using it don't repeat it
    from collections import Counter
    style_counts = Counter(json.dumps(e.get('s'), sort_keys=True) for e in cells.values() if e.get('s'))
    default_style = None
    if style_counts:
        key, count = style_counts.most_common(1)[0]
        if count >= 50:
            default_style = json.loads(key)
            # cells without a style in Excel must not pick up the default: give them an explicit empty one
            for r in range(1, last_r + 1):
                for c in range(1, last_c + 1):
                    coord = f'{openpyxl.utils.get_column_letter(c)}{r}'
                    e = cells.get(coord)
                    if e is None:
                        cells[coord] = {'s': {}}
                    elif not e.get('s'):
                        e['s'] = {}
            for coord in list(cells.keys()):
                e = cells[coord]
                cm = re.match(r'([A-Z]+)(\d+)$', coord)
                inside = int(cm.group(2)) <= last_r and openpyxl.utils.column_index_from_string(cm.group(1)) <= last_c
                if not inside:
                    continue  # the default style only applies inside the data area
                if e.get('s') is not None and json.dumps(e['s'], sort_keys=True) == key:
                    del e['s']
                    if 'v' not in e and 'f' not in e:
                        del cells[coord]

    sheet = {
        'maxRow': max(ws.max_row + 20, 100),
        'maxCol': max(ws.max_column + 2, 20),
        'colWidths': col_widths,
        'rowHeights': row_heights,
        'merges': [str(m) for m in ws.merged_cells.ranges],
        'cells': cells
    }
    if default_style:
        sheet['defaultStyle'] = default_style
        sheet['styleBounds'] = {'rows': last_r, 'cols': last_c}
    if ws.sheet_state != 'visible':
        sheet['hidden'] = ws.sheet_state  # 'hidden' | 'veryHidden'
    if hidden_cols: sheet['hiddenCols'] = sorted(set(hidden_cols), key=lambda x: (len(x), x))
    if hidden_rows: sheet['hiddenRows'] = sorted(set(hidden_rows))
    if freeze and (freeze['rows'] or freeze['cols']): sheet['freeze'] = freeze
    if filters:
        sheet['filters'] = filters
        sheet['filterHiddenRows'] = sorted(filtered)
    if IMAGES.get(ws.title):
        sheet['images'] = IMAGES[ws.title]
        print(f'  {ws.title}: imágenes={IMAGES[ws.title]}')
    if cond: sheet['condFormats'] = cond
    if ws.sheet_properties.tabColor is not None:
        tc = resolve_color(ws.sheet_properties.tabColor)
        if tc: sheet['tabColor'] = tc

    book['sheetOrder'].append(ws.title)
    book['sheets'][ws.title] = sheet
    stats[ws.title] = (len(cells), n_formulas, ws.sheet_state, freeze, len(cond))
    print(f'  {ws.title}: {len(cells)} celdas, {n_formulas} fórmulas, {ws.sheet_state}, freeze={freeze}, condicionales={len(cond)}', flush=True)

if PIVOTS:
    for i, p in enumerate(PIVOTS):
        p['id'] = f'pv{i + 1}'
    book['pivots'] = PIVOTS

active = wb.active.title if wb.active is not None else None
if active:
    book['activeSheet'] = active

with open(OUT, 'w', encoding='utf8') as fh:
    json.dump(book, fh, ensure_ascii=False, separators=(',', ':'))
print('OK ->', OUT)
